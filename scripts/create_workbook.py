"""Generate a blank journal; refuse to overwrite an existing user's workbook."""
import argparse
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties

HEADERS = ['법원','사건번호','물건번호','주소','유형','면적(㎡)','매각기일','감정가','최저가','유찰 횟수','예상 작성일','예상 낙찰가','가상 입찰가','판단 근거·특징','예상 시세','추가 비용 합계','목표 이익','내 입찰 상한가','상세 메모 링크','매각 결과','실제 낙찰가','입찰자 수','평가 가능','예측 오차율','절대 오차율','최저가 비교','실제 낙찰가 비교','상한가 비교','복기 메모','기록 점검']
GUIDE = [
('시작하기','파란 칸에 입력하고 회색 수식 칸은 유지하세요. 금액은 원, 면적은 ㎡, 날짜는 날짜 형식입니다. 300회차까지 준비되어 있습니다.'),
('한 행의 기준','법원·사건번호·물건번호·매각기일당 한 행입니다. 유찰 후 다음 기일은 새 행에 기록하세요. 사건번호는 연도를 포함하세요.'),
('세 가지 금액','예상 낙찰가=시장 낙찰가 예측 / 가상 입찰가=내 모의 입찰액 / 내 입찰 상한가=비용·목표 이익을 고려한 최대 지불액'),
('상한가','예상 시세 − 추가 비용 합계 − 목표 이익. 비용이나 목표 이익이 없으면 0을 명시하세요. 세 값이 모두 있어야 계산합니다.'),
('결과 전 기록','결과를 알기 전에 예상 작성일·예상 낙찰가·가상 입찰가·판단 근거를 기록하세요. 매각기일보다 늦은 예측은 평가에서 제외합니다. 같은 날은 결과 확인 전인지 직접 확인하세요.'),
('수정 전 값 보존','물건 기록의 원본 예측은 덮어쓰지 마세요. 수정은 예측 변경 이력에 이전값·변경값·변경 시각·사유를 추가하세요. 요약은 물건 기록의 원본 예측만 평가합니다. 파일 버전 기록도 함께 사용하세요. 자동 잠금·감사 기능은 없습니다.'),
('결과 입력','매각 결과를 선택하고 낙찰이면 실제 낙찰가와 입찰자 수를 기록하세요. 유찰·변경·연기·취하·미확인은 예측 평가에서 제외합니다.'),
('평가 대상','낙찰, 양수인 예상·실제 낙찰가, 예상 작성일≤매각기일, 식별값 완비, 중복 없음인 행만 집계합니다. 원본 예측의 사전 기록 여부는 사용자가 확인합니다.'),
('정확도','예측 오차율=(예상−실제)/실제. 양수는 과대예측, 음수는 과소예측입니다. 평균 절대 오차율이 낮을수록 정확합니다. ±5% 경계를 포함합니다.'),
('모의 입찰 비교','가상 입찰가를 최저가·실제 낙찰가·상한가와 각각 비교합니다. 실제 낙찰가와 동액은 별도 표시합니다. 초과라도 실제 낙찰을 보장하지 않으며 동액은 추첨 등 별도 판단이 필요합니다.'),
('수익성 분리','실제 낙찰가 초과 여부와 상한가 이내 여부를 별개로 확인하세요. 낙찰 가능성과 투자 수익성은 다른 판단입니다.'),
('비용과 메모','추가 비용은 취득세·등기·명도·수리·체납·금융·매도 비용 등을 상세 메모에서 산정하여 합계만 입력하세요. 메모 템플릿은 docs/property-note.md입니다.'),
('사용 환경','Excel에서 열거나 Google Drive에 업로드 후 Google 스프레드시트로 열어 사용하세요. 변환 후 수식·날짜·드롭다운을 확인하세요. 브라우저/Excel 실제 화면 검증은 별도로 필요합니다.'),
('20~30건 후','오차가 큰 원인과 자주 쓰는 메모를 검토한 뒤 전용 도구 필요성을 결정하세요. 실제 물건 데이터나 예시 거래는 포함하지 않았습니다.')]

def build(path, rows=300):
    if not isinstance(rows, int) or isinstance(rows, bool) or rows < 1:
        raise ValueError("rows must be a positive integer")
    last = rows + 1
    path = Path(path)
    if path.exists():
        raise FileExistsError(f'기존 파일을 덮어쓰지 않습니다: {path}')
    w=Workbook(); guide=w.active; guide.title='사용 안내'
    guide.append(['모의 입찰 기록장','madang · 원본 예측으로 시장 판단을 복기합니다'])
    for label, text in GUIDE: guide.append((label, text.replace("300회차", f"{rows}회차")))
    s=w.create_sheet('물건 기록'); s.append(HEADERS)
    for r in range(2,last+1):
        s.cell(r,18,f'=IF(COUNT(O{r}:Q{r})=3,O{r}-P{r}-Q{r},"")')
        key=f'COUNTIFS(A$2:A${last},A{r},B$2:B${last},B{r},C$2:C${last},C{r},G$2:G${last},G{r})'
        s.cell(r,23,f'=IF(AND(A{r}<>"",B{r}<>"",C{r}<>"",ISNUMBER(G{r}),ISNUMBER(K{r}),K{r}<=G{r},K{r}>0,T{r}="낙찰",ISNUMBER(L{r}),L{r}>0,ISNUMBER(U{r}),U{r}>0,{key}=1),"평가 가능","")')
        s.cell(r,24,f'=IF(W{r}="평가 가능",(L{r}-U{r})/U{r},"")')
        s.cell(r,25,f'=IF(W{r}="평가 가능",ABS(X{r}),"")')
        for c,ref,labels,guard in [(26,'I',['최저가 미달','최저가 동액','최저가 초과'],''),(27,'U',['실제 낙찰가 미만','실제 낙찰가 동액','실제 낙찰가 초과'],f'T{r}="낙찰",'),(28,'R',['상한가 미만','상한가 동액','상한가 초과'],'')]:
            s.cell(r,c,f'=IF(AND({guard}ISNUMBER(M{r}),M{r}>0,ISNUMBER({ref}{r})),IF(M{r}<{ref}{r},"{labels[0]}",IF(M{r}={ref}{r},"{labels[1]}","{labels[2]}")),"")')
        s.cell(r,30,f'=IF(COUNTA(A{r}:Q{r},S{r}:V{r},AC{r})=0,"",IF(OR(A{r}="",B{r}="",C{r}="",NOT(ISNUMBER(G{r}))),"식별값 확인",IF({key}>1,"중복 회차",IF(AND(ISNUMBER(K{r}),K{r}>G{r}),"기일 후 예측",IF(AND(T{r}="낙찰",W{r}=""),"평가 입력 확인","")))))')
    for c in [7,11]:
        for row in s.iter_rows(min_row=2,max_row=last,min_col=c,max_col=c): row[0].number_format='yyyy-mm-dd'
    for c in [8,9,12,13,15,16,17,18,21]:
        for row in s.iter_rows(min_row=2,max_row=last,min_col=c,max_col=c): row[0].number_format='#,##0;[Red]-#,##0'
    for row in s.iter_rows(min_row=2,max_row=last,min_col=24,max_col=25):
        for cell in row: cell.number_format='0.00%'
    for target,choices in [(f'E2:E{last}','아파트,다세대,단독주택,오피스텔,상가,토지,기타'),(f'T2:T{last}','미확인,낙찰,유찰,변경,연기,취하,기각')]:
        dv=DataValidation(type='list',formula1='"'+choices+'"',allow_blank=True); dv.errorTitle='목록에서 선택'; dv.error='허용된 값을 선택하세요.'; dv.showErrorMessage=True; s.add_data_validation(dv); dv.add(target)
    s.auto_filter.ref=f'A1:AD{last}'; s.freeze_panes='D2'
    summary=w.create_sheet('성과 요약'); summary.append(['지표','값'])
    for row in [('평가 가능한 건수','=COUNT(\'물건 기록\'!Y2:Y{last})'),('평균 절대 오차율','=IF(B2=0,"",AVERAGE(\'물건 기록\'!Y2:Y{last}))'),('평균 예측 오차율','=IF(B2=0,"",AVERAGE(\'물건 기록\'!X2:X{last}))'),('오차 5% 이내 건수','=COUNTIF(\'물건 기록\'!Y2:Y{last},"<=0.05")'),('오차 5% 이내 비율','=IF(B2=0,"",B5/B2)')]: summary.append(tuple(value.replace('{last}', str(last)) if isinstance(value, str) else value for value in row))
    for r in [3,4,6]: summary.cell(r,2).number_format='0.00%'
    summary.append(['해석','평균 절대 오차율은 정확도, 평균 예측 오차율은 과대·과소 편향입니다.'])
    history=w.create_sheet('예측 변경 이력'); history.append(['법원','사건번호','물건번호','매각기일','변경 시각','변경 항목','수정 전 값','수정 후 값','변경 사유','결과 확인 전 여부'])
    for _ in range(rows): history.append([None]*10)
    history.freeze_panes='E2'; history.auto_filter.ref=f'A1:J{last}'
    for ws in w:
        ws.sheet_view.showGridLines=False
        ws.row_dimensions[1].height=32
        for cell in ws[1]: cell.fill=PatternFill('solid',fgColor='17324D'); cell.font=Font(color='FFFFFF',bold=True); cell.alignment=Alignment(wrap_text=True,vertical='center')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                formula=cell.data_type=='f'; cell.font=Font(name='맑은 고딕',size=10,color='475569' if formula else '175A9B'); cell.fill=PatternFill('solid',fgColor='EEF2F6' if formula else ('F1F7FC' if cell.row%2==0 else 'FFFFFF')); cell.alignment=Alignment(vertical='top',wrap_text=True)
        for c in range(1,ws.max_column+1): ws.column_dimensions[col(c)].width=20
    guide.column_dimensions['A'].width=24; guide.column_dimensions['B'].width=110
    for r in range(2,guide.max_row+1): guide.row_dimensions[r].height=48
    for c in ['D','N','S','AC','AD']: s.column_dimensions[c].width=38
    summary.column_dimensions['A'].width=30; summary.column_dimensions['B'].width=80
    w.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True)
    path.parent.mkdir(parents=True,exist_ok=True); w.save(path)

if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('--output',default='templates/madang-auction-journal.xlsx'); p.add_argument('--rows',type=int,default=300); args=p.parse_args(); build(args.output, args.rows)
