import sys
import tempfile
import unittest
from pathlib import Path
import numpy as np
import formulas
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from scripts.create_workbook import build

class JournalTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory(); cls.path=Path(cls.tmp.name)/'journal.xlsx'; build(cls.path); cls.w=load_workbook(cls.path)
    def calc(self,cell,values):
        f=formulas.Parser().ast(self.w['물건 기록'][cell].value)[1].compile()
        values=values.copy()
        values['O2:Q2']=np.array([[values.get(c+'2','') for c in ['O','P','Q']]],dtype=object)
        return f(*[values.get(k,'') for k in f.inputs]).item()
    def test_upper_limit_and_blank(self):
        self.assertEqual(self.calc('R2',{'O2':200,'P2':30,'Q2':20}),150)
        self.assertEqual(self.calc('R2',{'O2':200,'P2':0,'Q2':0}),200)
        self.assertEqual(self.calc('R2',{'O2':200}), '')
    def test_comparisons(self):
        for bid,expected in [(90,'미달'),(100,'동액'),(110,'초과')]:
            self.assertEqual(self.calc('Z2',{'M2':bid,'I2':100}),'최저가 '+expected)
        self.assertEqual(self.calc('AA2',{'M2':100,'U2':100,'T2':'낙찰'}),'실제 낙찰가 동액')
        self.assertEqual(self.calc('AA2',{'M2':100,'U2':100,'T2':'유찰'}),'')
        self.assertEqual(self.calc('AB2',{'M2':110,'R2':100}),'상한가 초과')
    def test_errors_and_exclusion(self):
        self.assertAlmostEqual(self.calc('X2',{'W2':'평가 가능','L2':105,'U2':100}),.05)
        self.assertAlmostEqual(self.calc('X2',{'W2':'평가 가능','L2':95,'U2':100}),-.05)
        self.assertEqual(self.calc('Y2',{'W2':'평가 가능','X2':-.05}),.05)
        self.assertEqual(self.calc('X2',{'W2':'','U2':0}), '')
    def test_eligibility(self):
        v={'A2':'서울','B2':'2026타경1','C2':1,'G2':46000,'K2':45999,'T2':'낙찰','L2':100,'U2':100}
        for c in ['A','B','C','G']:
            v[f'{c}2:{c}301']=np.array([[v[f'{c}2']]]+[['']]*299,dtype=object)
        self.assertEqual(self.calc('W2',v),'평가 가능')
        for patch in [{'T2':'유찰'},{'K2':46001},{'U2':0},{'B2':''}]:
            self.assertEqual(self.calc('W2',v|patch),'')
        duplicate=v.copy(); duplicate['A2:A301']=np.array([['서울']]*300); duplicate['B2:B301']=np.array([['2026타경1']]*300); duplicate['C2:C301']=np.ones((300,1)); duplicate['G2:G301']=np.full((300,1),46000)
        self.assertEqual(self.calc('W2',duplicate),'')
    def test_summary_and_structure(self):
        ws=self.w['성과 요약']
        values={"'물건 기록'!Y2:Y301":np.array([[.05],[.15],['']],dtype=object),"'물건 기록'!X2:X301":np.array([[.05],[-.15],['']],dtype=object),'B2':2,'B5':1}
        for cell,expected in [('B2',2),('B3',.1),('B4',-.05),('B5',1),('B6',.5)]:
            f=formulas.Parser().ast(ws[cell].value)[1].compile()
            actual=np.asarray(f(*[values[k] for k in f.inputs])).item(); self.assertAlmostEqual(actual,expected)
        self.assertEqual(self.w['물건 기록'].max_row,301)
        self.assertEqual(len(self.w['물건 기록'].data_validations.dataValidation),2)
        with self.assertRaises(FileExistsError): build(self.path)

if __name__=='__main__': unittest.main()
